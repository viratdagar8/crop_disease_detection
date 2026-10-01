"""Generator for DOCX Documents:
1. Crop_Disease_Detection_SRS.docx (Software Requirements Specification)
2. Crop_Disease_Detection_Synopsis.docx (Official Project Synopsis)
3. Teacher_Explanation_Guide.docx (Sequential Explanation & Viva Study Guide)

Uses python-docx with clean academic formatting, professional tables, styling, and exact schema alignment.
"""

import sys
import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent

# Color Palette Constants
DARK_GREEN = RGBColor(27, 94, 32)      # #1B5E20
MED_GREEN = RGBColor(46, 125, 50)      # #2E7D32
CHARCOAL = RGBColor(33, 33, 33)        # #212121
ACCENT_ORANGE = RGBColor(230, 81, 0)   # #E65100
GRAY_BG = "F1F8E9"                     # Soft Green Background tint
BORDER_COLOR = "C8E6C9"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color_hex="C8E6C9"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def format_cell_padding(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = DARK_GREEN
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = MED_GREEN
    return h

def add_body_paragraph(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = CHARCOAL
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = CHARCOAL
    return p


# ==============================================================================
# 1. BUILD SOFTWARE REQUIREMENTS SPECIFICATION (SRS) DOCX
# ==============================================================================
def build_srs_docx():
    target_path = PROJECT_ROOT / "Crop_Disease_Detection_SRS.docx"
    doc = docx.Document()

    # Set Margins to 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # --------------------------------------------------------------------------
    # 1. Cover Page
    # --------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    r_t = p_title.add_run("SOFTWARE REQUIREMENTS SPECIFICATION (SRS)\n")
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_GREEN

    r_sub = p_title.add_run("Crop Disease Detection Using Convolutional Neural Networks (CNN)\nwith Hierarchical Binary Gate Verification")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = MED_GREEN

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(40)
    p_meta.paragraph_format.space_after = Pt(200)

    meta_text = (
        "Project Author: Virat Dagar\n"
        "Department: Computer Science & Engineering (Data Science)\n"
        "Institution: ABES Engineering College, Ghaziabad\n"
        "Program: HCLTech Industry Aligned Projects Program\n"
        "Academic Session: 2026–2027\n"
        "Document Version: 1.0 (Final Implemented SRS)"
    )
    r_m = p_meta.add_run(meta_text)
    r_m.font.name = 'Calibri'
    r_m.font.size = Pt(11)
    r_m.font.color.rgb = CHARCOAL

    doc.add_page_break()

    # --------------------------------------------------------------------------
    # 2. Abstract
    # --------------------------------------------------------------------------
    add_heading_1(doc, "2. Abstract")
    add_body_paragraph(
        doc,
        "Crop diseases cause significant financial destruction to farming communities and threaten global food security. "
        "Traditional manual scouting for crop pathology requires specialized agricultural expertise, which is scarce and expensive. "
        "This project delivers AgriScan AI, an automated Computer Vision decision-support application built with Deep Learning. "
        "The system incorporates a two-stage hierarchical model: Stage 1 (Approach 1: Binary Leaf Classifier Gate Model) filters non-leaf "
        "or out-of-focus background uploads, while Stage 2 evaluates valid plant leaves across 38 crop disease classes (14 crop species). "
        "Developed using Python 3.12, TensorFlow 2.18, and Flask 3.0, the system outputs instant disease diagnosis, empirical confidence scores, "
        "and actionable agronomic advisory guidelines. The expected outcome is early disease detection, reduced broad-spectrum chemical use, and protected crop yields."
    )

    # --------------------------------------------------------------------------
    # 3. Problem Statement
    # --------------------------------------------------------------------------
    add_heading_1(doc, "3. Problem Statement")
    add_body_paragraph(
        doc,
        "Agricultural disease scouting currently relies on visual inspection by agricultural extension officers or lab diagnostics. "
        "In rural regions, farmers face delayed diagnosis leading to rapid disease spread across entire fields. Furthermore, generic "
        "machine learning models frequently fail when non-leaf images are uploaded, outputting misleading predictions. There is a critical "
        "need for an automated, lightweight decision-support tool that validates input image quality via a binary gate model and accurately "
        "classifies crop leaf pathologies with integrated agronomic guidance."
    )

    # --------------------------------------------------------------------------
    # 4. Objectives
    # --------------------------------------------------------------------------
    add_heading_1(doc, "4. Objectives")
    objectives = [
        "1. Develop a user-friendly Flask web portal for uploading plant leaf photographs and viewing real-time diagnostics.",
        "2. Implement Approach 1: Binary Leaf Classifier (Gate Model) to reject invalid non-leaf images before disease evaluation.",
        "3. Train a 4-block Convolutional Neural Network (CNN) on 54,306 images across 38 plant disease/healthy classes.",
        "4. Integrate an agronomic knowledge base (disease_info.json) to deliver causes, symptoms, and organic/chemical treatments.",
        "5. Evaluate model performance using Categorical Accuracy, Weighted Precision, Recall, F1-score, and Confusion Matrix.",
        "6. Provide a production-ready, cloud-deployable web application compatible with Render."
    ]
    for obj in objectives:
        add_body_paragraph(doc, obj, space_after=2)

    # --------------------------------------------------------------------------
    # 5. Proposed Solution
    # --------------------------------------------------------------------------
    add_heading_1(doc, "5. Proposed Solution")
    add_body_paragraph(
        doc,
        "The proposed solution is AgriScan AI, a hierarchical multi-stage computer vision application. "
        "When a user uploads a photograph, the image is passed to the Stage 1 Binary Gate Model (Conv2D -> BatchNorm -> MaxPool -> Sigmoid) "
        "to verify if it represents a valid plant leaf. If passed, the image is normalized to [0,1] and fed to the Stage 2 Baseline CNN (38 classes). "
        "The top-1 prediction, top-3 candidates, confidence score, and agronomic management advice are rendered dynamically on the HTML5/JS web interface."
    )

    # --------------------------------------------------------------------------
    # 6. Scope of the Project
    # --------------------------------------------------------------------------
    add_heading_1(doc, "6. Scope of the Project")
    add_body_paragraph(
        doc,
        "The project scope encompasses 14 crop species (Apple, Blueberry, Cherry, Corn, Grape, Orange, Peach, Bell Pepper, Potato, Raspberry, Soybean, Squash, Strawberry, Tomato) "
        "covering 26 disease pathologies and 12 healthy controls (38 total classes). Implemented features include image format validation, Stage 1 binary gating, "
        "Stage 2 CNN inference, agronomic advisory output, and Flask web deployment. Out of current scope: live video stream processing and direct multi-label co-infection detection."
    )

    # --------------------------------------------------------------------------
    # 7. Functional Requirements
    # --------------------------------------------------------------------------
    add_heading_1(doc, "7. Functional Requirements")
    fr_data = [
        ["ID", "Functional Requirement", "Implementation Status"],
        ["FR-01", "User should be able to upload leaf images (.jpg, .png, .webp) via web interface.", "Implemented"],
        ["FR-02", "System shall validate file extension and file size limit (< 16 MB).", "Implemented"],
        ["FR-03", "System shall run Approach 1 Binary Gate Model to reject non-leaf uploads.", "Implemented"],
        ["FR-04", "System shall preprocess image (128x128 resizing, RGB normalization [0,1]).", "Implemented"],
        ["FR-05", "System shall execute Stage 2 38-class CNN inference and generate Softmax probabilities.", "Implemented"],
        ["FR-06", "System shall display top-1 headline prediction, confidence score %, and top-3 candidates.", "Implemented"],
        ["FR-07", "System shall fetch and display agronomic advisory (symptoms, cause, treatment) from knowledge base.", "Implemented"],
        ["FR-08", "System shall provide health check (/health) and supported classes (/classes) REST endpoints.", "Implemented"]
    ]
    t_fr = doc.add_table(rows=len(fr_data), cols=3)
    t_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_fr)
    for r_idx, row in enumerate(fr_data):
        for c_idx, val in enumerate(row):
            cell = t_fr.cell(r_idx, c_idx)
            cell.text = val
            format_cell_padding(cell)
            if r_idx == 0:
                set_cell_background(cell, GRAY_BG)
                p = cell.paragraphs[0]
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = DARK_GREEN

    # --------------------------------------------------------------------------
    # 8. Non-Functional Requirements
    # --------------------------------------------------------------------------
    add_heading_1(doc, "8. Non-Functional Requirements")
    nfrs = [
        ("Performance: ", "Single-image inference latency is under 250 milliseconds on standard CPU hardware."),
        ("Usability: ", "Clean, responsive drag-and-drop web portal suitable for non-technical agricultural workers."),
        ("Reliability: ", "Graceful error handling for corrupted image files or non-leaf uploads via Stage 1 gating."),
        ("Compatibility: ", "Fully compatible with all modern desktop and mobile browsers (Chrome, Edge, Firefox, Safari)."),
        ("Security: ", "Client-side image payload validation preventing unauthorized file script execution.")
    ]
    for n_title, n_desc in nfrs:
        add_body_paragraph(doc, n_desc, bold_prefix=n_title, space_after=3)

    # --------------------------------------------------------------------------
    # 9. Technology Stack
    # --------------------------------------------------------------------------
    add_heading_1(doc, "9. Technology Stack")
    tech_data = [
        ["Component", "Technology Used", "Version / Detail"],
        ["Programming Language", "Python", "3.12.10"],
        ["Deep Learning Framework", "TensorFlow / Keras", "2.18.0 / 3.15.1"],
        ["Backend Web Framework", "Flask", "3.0.0 (Gunicorn WSGI)"],
        ["Frontend Interface", "HTML5, CSS3, JavaScript (ES6)", "Responsive DOM UI"],
        ["Computer Vision / Image Processing", "Pillow (PIL), NumPy", "10.0.0+ / 1.24.0+"],
        ["Dataset Source", "Hugging Face Datasets", "mohanty/PlantVillage (color)"],
        ["PDF & Document Generation", "ReportLab, python-docx", "4.0.0+ / 1.2.0+"],
        ["Version Control & Cloud Hosting", "Git, Render", "Render Cloud Web Service"]
    ]
    t_tech = doc.add_table(rows=len(tech_data), cols=3)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech)
    for r_idx, row in enumerate(tech_data):
        for c_idx, val in enumerate(row):
            cell = t_tech.cell(r_idx, c_idx)
            cell.text = val
            format_cell_padding(cell)
            if r_idx == 0:
                set_cell_background(cell, GRAY_BG)
                p = cell.paragraphs[0]
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = DARK_GREEN

    # --------------------------------------------------------------------------
    # 10. System Architecture
    # --------------------------------------------------------------------------
    add_heading_1(doc, "10. System Architecture")
    add_body_paragraph(
        doc,
        "The system follows a modular 3-tier architectural flow:\n"
        "[ USER ] --> [ FRONTEND (HTML/JS) ] --> [ BACKEND API (app.py) ]\n"
        "                                               │\n"
        "                                               ▼\n"
        "                                 [ STAGE 1: GATE MODEL (CNN) ]\n"
        "                                       │              │\n"
        "                            (Non-Leaf) │              │ (Valid Leaf)\n"
        "                                       ▼              ▼\n"
        "                              [ REJECT WARNING ]   [ STAGE 2: 38-CLASS CNN ]\n"
        "                                                      │\n"
        "                                                      ▼\n"
        "                                           [ AGRONOMIC ADVISORY ]"
    )

    # --------------------------------------------------------------------------
    # 11. System Design
    # --------------------------------------------------------------------------
    add_heading_1(doc, "11. System Design")
    add_heading_2(doc, "11.1 Use Case & Data Flow")
    add_body_paragraph(
        doc,
        "1. User selects/drags a leaf image on the web interface.\n"
        "2. Client JavaScript validates extension (.jpg/.png) and posts multipart payload to /predict.\n"
        "3. Flask backend converts image to RGB and triggers src/predict.py.\n"
        "4. Stage 1 Gate Model classifies Leaf vs Non-Leaf.\n"
        "5. Stage 2 baseline CNN computes Softmax probabilities across 38 classes.\n"
        "6. Diagnostic payload mapped with disease_info.json advisory is returned as JSON to frontend."
    )
    add_heading_2(doc, "11.2 CNN Architecture Design")
    add_body_paragraph(
        doc,
        "• Block 1: Conv2D(32, 3x3) -> BatchNorm -> MaxPool(2x2) -> Dropout(0.2)\n"
        "• Block 2: Conv2D(64, 3x3) -> BatchNorm -> MaxPool(2x2) -> Dropout(0.2)\n"
        "• Block 3: Conv2D(128, 3x3) -> BatchNorm -> MaxPool(2x2) -> Dropout(0.25)\n"
        "• Block 4: Conv2D(256, 3x3) -> BatchNorm -> MaxPool(2x2) -> Dropout(0.25)\n"
        "• Classifier: Flatten -> Dense(256) -> BatchNorm -> Dropout(0.5) -> Dense(38, Softmax)"
    )

    # --------------------------------------------------------------------------
    # 12. Module Description
    # --------------------------------------------------------------------------
    add_heading_1(doc, "12. Module Description")
    modules = [
        ("Module 1 - Image Ingestion & Input Validation: ", "Handles upload staging, file extension verification, PIL RGB conversion, and 16 MB payload capping (app.py)."),
        ("Module 2 - Approach 1 Binary Leaf Gate Model: ", "Stage 1 CNN model that evaluates chlorophyll chromaticity and binary leaf features to reject non-leaf uploads (src/gate_model.py)."),
        ("Module 3 - Stage 2 Multiclass CNN Classifier: ", "Core 4-block neural network computing class logits across 38 plant disease/healthy categories (src/model.py & src/predict.py)."),
        ("Module 4 - Agronomic Advisory Engine: ", "JSON knowledge base containing causes, symptoms, and organic/chemical management guidelines for all 38 classes (model/disease_info.json)."),
        ("Module 5 - Flask Web Portal: ", "REST API endpoints (/predict, /health, /classes) and responsive single-page user interface (templates/index.html & static/).")
    ]
    for m_title, m_desc in modules:
        add_body_paragraph(doc, m_desc, bold_prefix=m_title, space_after=4)

    # --------------------------------------------------------------------------
    # 13. Metadata & Knowledge Base Design
    # --------------------------------------------------------------------------
    add_heading_1(doc, "13. Metadata & Knowledge Base Design")
    add_body_paragraph(doc, "Structure of class_indices.json and disease_info.json metadata repositories:")
    meta_data = [
        ["Key Field", "Data Type", "Description / Sample Value"],
        ["class_id", "Integer (0..37)", "Class index mapped to folder label (e.g. 0: 'Apple___Apple_scab')"],
        ["crop", "String", "Clean crop species name (e.g. 'Apple', 'Tomato')"],
        ["disease", "String", "Pathology name (e.g. 'Early Blight', 'Healthy Foliage')"],
        ["symptoms", "String", "Detailed visual symptoms exhibited on leaf blade"],
        ["management", "String", "Actionable cultural, organic, and chemical control advice"]
    ]
    t_m = doc.add_table(rows=len(meta_data), cols=3)
    t_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_m)
    for r_idx, row in enumerate(meta_data):
        for c_idx, val in enumerate(row):
            cell = t_m.cell(r_idx, c_idx)
            cell.text = val
            format_cell_padding(cell)
            if r_idx == 0:
                set_cell_background(cell, GRAY_BG)
                p = cell.paragraphs[0]
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = DARK_GREEN

    # --------------------------------------------------------------------------
    # 14. Testing
    # --------------------------------------------------------------------------
    add_heading_1(doc, "14. Testing")
    test_data = [
        ["Test ID", "Test Case Description", "Expected Result", "Status"],
        ["TC01", "Upload valid leaf image (.jpg)", "Displays headline prediction, confidence, top-3 candidates, and advisory.", "PASS"],
        ["TC02", "Upload non-leaf image (e.g. background artifact)", "Stage 1 Gate Model rejects image with warning message.", "PASS"],
        ["TC03", "Upload invalid file format (.txt/.exe)", "Client/Server validation returns HTTP 400 bad request error.", "PASS"],
        ["TC04", "Upload oversized file (> 16 MB)", "Server triggers HTTP 413 Payload Too Large error.", "PASS"],
        ["TC05", "GET /health REST endpoint", "Returns JSON status: healthy with model_loaded: true.", "PASS"],
        ["TC06", "GET /classes REST endpoint", "Returns complete list of 38 supported crop disease classes.", "PASS"]
    ]
    t_test = doc.add_table(rows=len(test_data), cols=4)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_test)
    for r_idx, row in enumerate(test_data):
        for c_idx, val in enumerate(row):
            cell = t_test.cell(r_idx, c_idx)
            cell.text = val
            format_cell_padding(cell)
            if r_idx == 0:
                set_cell_background(cell, GRAY_BG)
                p = cell.paragraphs[0]
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = DARK_GREEN

    # --------------------------------------------------------------------------
    # 15. Limitations
    # --------------------------------------------------------------------------
    add_heading_1(doc, "15. Limitations")
    limits = [
        "1. Single Infection Assumption: Assumes one primary disease per leaf photograph.",
        "2. Controlled Background Sensitivity: PlantVillage images were captured under lab lighting; complex outdoor backgrounds require Stage 1 gating.",
        "3. CPU Inference Capping: Batch evaluation relies on CPU execution without dedicated local GPU hardware acceleration.",
        "4. Non-Medical Advisory: Advisory serves as a decision support tool and does not replace lab chemical analysis."
    ]
    for lim in limits:
        add_body_paragraph(doc, lim, space_after=2)

    # --------------------------------------------------------------------------
    # 16. Future Scope
    # --------------------------------------------------------------------------
    add_heading_1(doc, "16. Future Scope")
    futures = [
        "• Native Mobile App: Flutter/Android application with offline TFLite model inference for remote field usage.",
        "• Multi-Label Segmentation: Mask R-CNN integration for segmenting multiple co-occurring diseases on single leaves.",
        "• IoT Drone Integration: Aerial crop scouting via automated drone imagery streaming.",
        "• Multilingual Support: Local language translation (Hindi, regional dialects) for agricultural advisory."
    ]
    for fut in futures:
        add_body_paragraph(doc, fut, space_after=2)

    # --------------------------------------------------------------------------
    # 17. References
    # --------------------------------------------------------------------------
    add_heading_1(doc, "17. References")
    refs = [
        "1. Mohanty, S. P., Hughes, D. P., & Salathé, M. (2016). Using Deep Learning for Image-Based Plant Disease Detection. Frontiers in Plant Science, 7, 1419.",
        "2. TensorFlow & Keras Official Documentation: https://www.tensorflow.org/",
        "3. Hugging Face PlantVillage Dataset Repository: https://huggingface.co/datasets/mohanty/PlantVillage",
        "4. Flask Web Development Documentation: https://flask.palletsprojects.com/",
        "5. Render Cloud Deployment Documentation: https://render.com/docs"
    ]
    for ref in refs:
        add_body_paragraph(doc, ref, space_after=2)

    doc.save(str(target_path))
    print(f"[OK] Generated SRS DOCX: {target_path}")


# ==============================================================================
# 2. BUILD PROJECT SYNOPSIS DOCX
# ==============================================================================
def build_synopsis_docx():
    target_path = PROJECT_ROOT / "Crop_Disease_Detection_Synopsis.docx"
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title / Header
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(12)
    p_t.paragraph_format.space_after = Pt(6)
    r = p_t.add_run("PROJECT SYNOPSIS\n")
    r.font.name = 'Calibri'
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = DARK_GREEN

    r2 = p_t.add_run("Crop Disease Detection Using Convolutional Neural Networks (CNN)")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = MED_GREEN

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(18)
    p_meta.add_run(
        "Student: Virat Dagar | Dept: Computer Science & Engineering (Data Science)\n"
        "Institution: ABES Engineering College, Ghaziabad\n"
        "Program: HCLTech Industry Aligned Projects (2026–2027)"
    )

    # 1. Project Title
    add_heading_1(doc, "1. Project Title")
    add_body_paragraph(doc, "Crop Disease Detection Using Convolutional Neural Networks (CNN) with Hierarchical Binary Gate Verification.")

    # 2. Introduction
    add_heading_1(doc, "2. Introduction")
    add_body_paragraph(
        doc,
        "Plant diseases pose a severe threat to global agricultural yield, causing significant economic loss and endangering food security. "
        "Traditional disease diagnosis relies on manual visual scouting by agricultural extension experts, which is labor-intensive and unavailable to rural farmers. "
        "This project introduces AgriScan AI, a deep learning decision-support system that automatically classifies 38 plant leaf pathologies across 14 crop species."
    )

    # 3. Problem Statement
    add_heading_1(doc, "3. Problem Statement")
    add_body_paragraph(
        doc,
        "Delayed identification of crop pathogens leads to unchecked disease spread and excessive chemical application. "
        "Furthermore, existing machine learning classifiers output erroneous disease predictions when non-leaf images are uploaded. "
        "There is an immediate need for an automated tool with binary gating verification to filter bad inputs and deliver reliable disease diagnostics."
    )

    # 4. Objectives
    add_heading_1(doc, "4. Objectives")
    objs = [
        "1. Develop an accessible web application for uploading crop leaf images.",
        "2. Implement Approach 1 Binary Gate Model to filter out non-leaf images.",
        "3. Train a 4-block CNN model on 54,306 images across 38 crop disease classes.",
        "4. Provide detailed agronomic advisory guidelines for disease prevention and treatment.",
        "5. Deploy the system to cloud infrastructure for real-time demonstration."
    ]
    for o in objs:
        add_body_paragraph(doc, o, space_after=2)

    # 5. Proposed Solution
    add_heading_1(doc, "5. Proposed Solution")
    add_body_paragraph(
        doc,
        "AgriScan AI uses a hierarchical multi-stage model pipeline. Stage 1 verifies leaf validity using Approach 1 Binary Gate Classifier. "
        "Stage 2 computes class probabilities across 38 target categories using a baseline CNN. Stage 3 maps predictions to an agronomic knowledge base."
    )

    # 6. Scope of the Project
    add_heading_1(doc, "6. Scope of the Project")
    add_body_paragraph(
        doc,
        "Covers 14 crop species (Apple, Cherry, Corn, Grape, Peach, Pepper, Potato, Strawberry, Tomato, etc.) and 38 total classes. "
        "Includes web upload UI, binary gating, CNN classification, advisory retrieval, and Cloud WSGI deployment."
    )

    # 7. Methodology / Working
    add_heading_1(doc, "7. Methodology / Working")
    add_body_paragraph(
        doc,
        "DATASET (PlantVillage 54,306 images) --> PREPROCESSING (128x128 Resizing & Normalization) --> "
        "STAGE 1 (Approach 1 Binary Gate Model) --> STAGE 2 (38-Class CNN Classifier) --> "
        "ADVISORY MAPPING (disease_info.json) --> FLASK WEB UI OUTPUT"
    )

    # 8. Technologies Used
    add_heading_1(doc, "8. Technologies Used")
    t_tech = doc.add_table(rows=6, cols=2)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech)
    data = [
        ["Category", "Technology"],
        ["Programming Language", "Python 3.12"],
        ["Deep Learning", "TensorFlow 2.18 / Keras"],
        ["Backend / Web", "Flask 3.0 (Gunicorn)"],
        ["Frontend", "HTML5, CSS3, JavaScript"],
        ["Dataset", "Hugging Face PlantVillage"]
    ]
    for r_i, row in enumerate(data):
        for c_i, val in enumerate(row):
            cell = t_tech.cell(r_i, c_i)
            cell.text = val
            format_cell_padding(cell)
            if r_i == 0:
                set_cell_background(cell, GRAY_BG)
                cell.paragraphs[0].runs[0].font.bold = True

    # 9. Modules
    add_heading_1(doc, "9. Modules")
    mods = [
        ("Module 1 - Ingestion: ", "File upload, format validation, and image normalization."),
        ("Module 2 - Approach 1 Gate: ", "Binary CNN verifying leaf presence."),
        ("Module 3 - CNN Classifier: ", "38-class disease categorization network."),
        ("Module 4 - Advisory Engine: ", "JSON knowledge base for agronomic treatment."),
        ("Module 5 - Web Dashboard: ", "Flask web portal and REST API endpoints.")
    ]
    for m_t, m_d in mods:
        add_body_paragraph(doc, m_d, bold_prefix=m_t, space_after=3)

    # 10. Expected Outcome
    add_heading_1(doc, "10. Expected Outcome")
    add_body_paragraph(doc, "Instant, highly accurate crop disease diagnosis with actionable treatment advice, reducing manual scouting costs and protecting crop yields.")

    # 11. Future Scope
    add_heading_1(doc, "11. Future Scope")
    add_body_paragraph(doc, "Offline mobile app deployment via TFLite, multi-label co-infection detection, drone imagery integration, and regional language support.")

    # 12. References
    add_heading_1(doc, "12. References")
    add_body_paragraph(doc, "1. Mohanty et al. (2016), Frontiers in Plant Science.\n2. TensorFlow & Keras Official Documentation.\n3. PlantVillage Dataset on Hugging Face.")

    doc.save(str(target_path))
    print(f"[OK] Generated Synopsis DOCX: {target_path}")


# ==============================================================================
# 3. BUILD TEACHER EXPLANATION & VIVA STUDY GUIDE DOCX
# ==============================================================================
def build_teacher_guide_docx():
    target_path = PROJECT_ROOT / "Teacher_Explanation_Guide.docx"
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(12)
    p_t.paragraph_format.space_after = Pt(6)
    r = p_t.add_run("TEACHER EXPLANATION & VIVA DEFENSE GUIDE\n")
    r.font.name = 'Calibri'
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = DARK_GREEN

    r2 = p_t.add_run("Step-by-Step Technical Masterclass for Project Evaluation")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = MED_GREEN

    add_heading_1(doc, "1. How to Present This Project to Your Teacher (Sequential Explanation)")
    steps = [
        ("Step 1: Start with the Real-World Problem Statement",
         "Tell your teacher: 'Respected Ma'am/Sir, crop diseases cause up to 40% yield loss globally. Smallholder farmers lack immediate access to plant pathologists. Our project, AgriScan AI, provides an automated decision-support tool using Deep Learning.'"),

        ("Step 2: Highlight Approach 1 (Binary Leaf Gate Model)",
         "Explain: 'Most standard ML models fail when a user uploads a non-leaf image (like a car or shoe) by making an arbitrary disease guess. To solve this, we implemented Approach 1: Binary Leaf Classifier (Gate Model). It acts as a Stage 1 security gate, checking if P(Is Leaf) >= 0.5 before allowing the image into the disease classifier.'"),

        ("Step 3: Explain the Stage 2 CNN Architecture",
         "Explain: 'For valid leaves, we pass a 128x128x3 normalized image tensor into a 4-block Baseline CNN. Block 1 detects low-level edges, Block 2 extracts leaf veins, Block 3 detects lesion spots, and Block 4 learns disease patterns. The Softmax layer outputs probabilities across 38 target classes.'"),

        ("Step 4: Demonstrate the Web Interface & Advisory Engine",
         "Explain: 'We built a Flask 3.0 backend with a responsive drag-and-drop web UI. When a leaf is uploaded, the app outputs the top-1 headline prediction, confidence %, top-3 candidates, and retrieves treatment advice (cause, symptoms, organic/chemical management) from disease_info.json.'")
    ]
    for s_title, s_desc in steps:
        add_body_paragraph(doc, s_desc, bold_prefix=s_title + "\n", space_after=6)

    add_heading_1(doc, "2. Top Teacher Viva Voce Questions & High-Score Answers")
    qas = [
        ("Q1: Why did you choose CNN over traditional ML algorithms like SVM or Random Forest?",
         "Answer: Traditional ML requires manual feature extraction (e.g. SIFT, color histograms) which cannot capture subtle pathological variations across 38 classes. CNNs automatically learn hierarchical visual features directly from raw pixels (Edges -> Textures -> Lesions -> Disease Patterns) with spatial and translation invariance."),

        ("Q2: Why use SparseCategoricalCrossentropy instead of CategoricalCrossentropy?",
         "Answer: CategoricalCrossentropy requires one-hot encoded ground truth vectors for 38 classes, consuming excessive memory over 54,000 images. SparseCategoricalCrossentropy accepts integer labels directly, achieving identical mathematical loss calculation with significantly lower memory overhead."),

        ("Q3: How do Batch Normalization and Dropout work in your model?",
         "Answer: Batch Normalization normalizes layer activations across mini-batches, accelerating training convergence. Dropout randomly zeroes 20% to 50% of neuron outputs during forward passes, preventing feature co-adaptation and strong overfitting on leaf background textures."),

        ("Q4: What dataset did you use and what is its split structure?",
         "Answer: We used the official PlantVillage dataset (color configuration) hosted on Hugging Face (mohanty/PlantVillage), containing 54,306 RGB images across 14 crop species and 38 classes. We used an 80/20 train/test split with 15% validation carving for model checkpointing.")
    ]
    for q, a in qas:
        p_q = doc.add_paragraph()
        r_q = p_q.add_run(q)
        r_q.font.bold = True
        r_q.font.color.rgb = ACCENT_ORANGE
        add_body_paragraph(doc, a, space_after=6)

    doc.save(str(target_path))
    print(f"[OK] Generated Teacher Guide DOCX: {target_path}")


if __name__ == "__main__":
    print("Generating DOCX Deliverables...")
    build_srs_docx()
    build_synopsis_docx()
    build_teacher_guide_docx()
    print("[OK] All DOCX files generated successfully!")
